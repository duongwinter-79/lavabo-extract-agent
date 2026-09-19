"""Phone photos and demo videos -> the images/ folder the pack expects.

Nobody renames files on a phone. Asking a shop owner to produce
`BC52-80-TRANG__front.jpg` from an iPhone is asking for nothing, politely, so the pack
would arrive with no photos and the agent would answer "cho em xem mẫu" with words.

What a phone CAN do easily is make a folder per product in the Drive app and drop photos
into it. So that is the input: one folder per mã SP, any filenames inside. This renames,
downscales, pulls frames out of demo videos, and writes the images.xlsx mapping.

    data-tu-dien-thoai/
      BC52-80-TRANG/  IMG_4821.HEIC  IMG_4822.jpg  demo.mp4
      GUONG-BO-60/    IMG_5001.jpg

A file sitting loose at the top is reported rather than guessed at: an unattached photo
costs a missing picture, a misattached one puts the wrong product in front of a customer.
"""

from __future__ import annotations

import logging
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from ..video import ffmpeg_path
from .spec import IMAGES

log = logging.getLogger(__name__)

PHOTO_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}
# Apple's default since iOS 11. Pillow cannot read it without pillow-heif, and telling
# the shop to flip one iPhone setting is cheaper than adding a dependency for it.
HEIC_SUFFIXES = {".heic", ".heif"}
VIDEO_SUFFIXES = {".mp4", ".mov", ".m4v", ".avi", ".3gp"}

# Roles in the order a product is usually photographed, so an unhinted batch still lands
# on sensible names.
ROLES = ("front", "angle", "detail", "lapdat", "size")
ROLE_HINTS = {
    # No bare digits here. "1" once matched the frame named "video1" and quietly
    # overwrote the frame already saved as front -- a hint has to be a word.
    "front": ("front", "chinh", "thang"),
    "angle": ("angle", "nghieng", "goc"),
    "detail": ("detail", "chitiet", "canh", "close"),
    "lapdat": ("lapdat", "lap", "thucte", "install"),
    "size": ("size", "kichthuoc", "banve", "spec"),
}

MAX_EDGE = 1600
JPEG_QUALITY = 85
FRAMES_PER_VIDEO = 3


@dataclass
class MediaReport:
    products: dict[str, list[str]] = field(default_factory=dict)
    videos: list[str] = field(default_factory=list)
    heic: list[str] = field(default_factory=list)
    unmapped: list[str] = field(default_factory=list)
    unknown_sku: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)

    @property
    def photo_count(self) -> int:
        return sum(len(v) for v in self.products.values())


def organise_workbook(path: Path, images_dir: Path, *,
                      known_skus: set[str] | None = None) -> MediaReport:
    """Pull images out of a workbook somebody pasted them into.

    Telling a shop "put photos in a folder, not in the spreadsheet" is correct and will
    be ignored by about half of them, because pasting a picture next to the product is
    the obvious thing to do. openpyxl keeps the anchor, so the row a picture sits on
    identifies the product, and the intuitive act stops being a dead end.

    What this cannot see: images inserted *into a cell* by Google Sheets, or an =IMAGE()
    formula. Those are not drawings and never reach the file as one. **[confirm]** against
    a real Sheets export; when in doubt the folder route always works.
    """
    from openpyxl import load_workbook

    report = MediaReport()
    images_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(path)

    for ws in wb.worksheets:
        header = [str(c).strip().lower() if c is not None else ""
                  for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        if "ma_sp" not in header:
            continue
        column = header.index("ma_sp")

        by_row: dict[int, str] = {}
        for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if column < len(values) and values[column]:
                by_row[number] = str(values[column]).strip()

        for image in getattr(ws, "_images", []):
            row = getattr(image.anchor, "_from", None)
            sku = by_row.get(row.row + 1) if row is not None else None
            if not sku:
                report.unmapped.append(f"{ws.title}: ảnh không nằm trên dòng sản phẩm nào")
                continue
            saved = report.products.setdefault(sku, [])
            name = _target_name(sku, "", len(saved), saved, use_hints=False)
            if _write_bytes(image._data(), images_dir / name):
                saved.append(name)
            else:
                report.skipped.append(f"{sku}: một ảnh không đọc được")

    if known_skus is not None:
        report.unknown_sku = sorted(s for s in report.products if s.lower() not in known_skus)
    return report


def _write_bytes(blob: bytes, target: Path) -> bool:
    """Downscale from memory, so an embedded photo gets the same treatment as a file."""
    import io

    try:
        from PIL import Image, ImageOps
    except ImportError:
        target.write_bytes(blob)
        return True
    try:
        with Image.open(io.BytesIO(blob)) as img:
            img = ImageOps.exif_transpose(img).convert("RGB")
            img.thumbnail((MAX_EDGE, MAX_EDGE))
            img.save(target, "JPEG", quality=JPEG_QUALITY, optimize=True)
        return True
    except Exception as exc:
        log.warning("bỏ qua một ảnh nhúng: %s", exc)
        return False


def organise(source: Path, images_dir: Path, *, known_skus: set[str] | None = None,
             frames: int = FRAMES_PER_VIDEO,
             mapping: dict[str, tuple[str, str]] | None = None) -> MediaReport:
    report = MediaReport()
    images_dir.mkdir(parents=True, exist_ok=True)

    if mapping is not None:
        return _from_mapping(source, images_dir, mapping, known_skus, report, frames)

    for entry in sorted(source.iterdir()):
        if entry.is_dir():
            _product_folder(entry, images_dir, report, frames)
        elif entry.is_file():
            if sku := _sku_from_name(entry.name):
                _one_file(entry, sku, images_dir, report, frames)
            elif entry.suffix.lower() in PHOTO_SUFFIXES | HEIC_SUFFIXES | VIDEO_SUFFIXES:
                report.unmapped.append(entry.name)

    if known_skus is not None:
        report.unknown_sku = sorted(s for s in report.products if s.lower() not in known_skus)
    return report


def _from_mapping(source: Path, images_dir: Path, mapping: dict[str, tuple[str, str]],
                  known_skus: set[str] | None, report: MediaReport,
                  frames: int) -> MediaReport:
    """Name photos from a filled contact sheet rather than from their folder.

    A file with no row is left unmapped rather than filed somewhere plausible: the whole
    reason this path exists is that nobody could tell which product it was.
    """
    for path in sorted(p for p in source.rglob("*") if p.is_file()):
        if path.suffix.lower() not in PHOTO_SUFFIXES | HEIC_SUFFIXES | VIDEO_SUFFIXES:
            continue
        entry = mapping.get(path.name.lower())
        if not entry:
            report.unmapped.append(path.name)
            continue
        sku, role = entry
        saved = report.products.setdefault(sku, [])
        if path.suffix.lower() in HEIC_SUFFIXES:
            report.heic.append(f"{sku}/{path.name}")
            continue
        if path.suffix.lower() in VIDEO_SUFFIXES:
            report.videos.append(f"{sku}/{path.name}")
            saved.extend(_video_frames(path, sku, images_dir, frames, len(saved), saved))
            continue
        name = _target_name(sku, role or "", len(saved), saved, use_hints=bool(role))
        if _downscale(path, images_dir / name):
            saved.append(name)
        else:
            report.skipped.append(path.name)

    if known_skus is not None:
        report.unknown_sku = sorted(s for s in report.products if s.lower() not in known_skus)
    return report


def _product_folder(folder: Path, images_dir: Path, report: MediaReport, frames: int) -> None:
    sku = folder.name.strip()
    files = sorted(p for p in folder.rglob("*") if p.is_file())
    for path in files:
        _one_file(path, sku, images_dir, report, frames)
    report.products.setdefault(sku, [])


def _one_file(path: Path, sku: str, images_dir: Path, report: MediaReport,
              frames: int) -> None:
    suffix = path.suffix.lower()
    saved = report.products.setdefault(sku, [])

    if suffix in HEIC_SUFFIXES:
        report.heic.append(f"{sku}/{path.name}")
        return
    if suffix in VIDEO_SUFFIXES:
        report.videos.append(f"{sku}/{path.name}")
        saved.extend(_video_frames(path, sku, images_dir, frames, len(saved), saved))
        return
    if suffix not in PHOTO_SUFFIXES:
        report.skipped.append(path.name)
        return

    name = _target_name(sku, path.name, len(saved), saved)
    if _downscale(path, images_dir / name):
        saved.append(name)
    else:
        report.skipped.append(path.name)


def _target_name(sku: str, original: str, index: int, taken: list[str],
                 *, use_hints: bool = True) -> str:
    """A name for this photo that is not already in use for this product.

    Two photos of the same product can easily both look like a "front" to the hints, and
    the loser of that race used to be overwritten on disk without a word.
    """
    if use_hints:
        stem = re.sub(r"[^a-z0-9]", "", original.lower())
        for role, hints in ROLE_HINTS.items():
            candidate = f"{sku}__{role}.jpg"
            if any(h in stem for h in hints) and candidate not in taken:
                return candidate

    position = index
    while True:
        role = (ROLES[position] if position < len(ROLES)
                else f"extra{position - len(ROLES) + 1}")
        candidate = f"{sku}__{role}.jpg"
        if candidate not in taken:
            return candidate
        position += 1


def _downscale(source: Path, target: Path) -> bool:
    """Phone photos are 3-6MB; Messenger and Meta's catalogue both want far less."""
    try:
        from PIL import Image, ImageOps
    except ImportError:
        shutil.copyfile(source, target)
        return True

    try:
        with Image.open(source) as img:
            img = ImageOps.exif_transpose(img)     # a photo taken sideways stays upright
            img = img.convert("RGB")
            img.thumbnail((MAX_EDGE, MAX_EDGE))
            img.save(target, "JPEG", quality=JPEG_QUALITY, optimize=True)
        return True
    except Exception as exc:                        # a truncated download, a fake .jpg
        log.warning("bỏ qua %s: %s", source.name, exc)
        return False


def _video_frames(path: Path, sku: str, images_dir: Path, frames: int,
                  start_index: int, taken: list[str]) -> list[str]:
    """Candidate stills from a demo video, sharpest first.

    A demo video is not a catalogue photo, but it is usually the only picture of the
    product actually installed, and the shop already has it. Picking by sharpness keeps
    the motion-blurred frames out.
    """
    try:
        ffmpeg = ffmpeg_path()
    except Exception as exc:
        log.warning("không đọc được video %s: %s", path.name, exc)
        return []

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        proc = subprocess.run(
            [ffmpeg, "-loglevel", "error", "-i", str(path), "-vf", "fps=1",
             "-frames:v", "30", str(out / "f%03d.jpg")],
            capture_output=True)
        if proc.returncode != 0:
            log.warning("ffmpeg không đọc được %s", path.name)
            return []

        candidates = sorted(out.glob("*.jpg"))
        if not candidates:
            return []
        best = _sharpest(candidates, frames)

        names: list[str] = []
        for offset, frame in enumerate(best):
            # Frame filenames carry no information about the shot, so never hint off them.
            name = _target_name(sku, "", start_index + offset, taken + names,
                                use_hints=False)
            if _downscale(frame, images_dir / name):
                names.append(name)
        return names


def _sharpest(paths: list[Path], want: int) -> list[Path]:
    try:
        import numpy as np
        from PIL import Image
    except ImportError:
        step = max(len(paths) // max(want, 1), 1)
        return paths[::step][:want]

    scored: list[tuple[float, Path]] = []
    for path in paths:
        with Image.open(path) as img:
            grey = np.asarray(img.convert("L"), dtype=np.float32)
        # Same measure video.py uses to tell a still frame from a smeared one.
        scored.append((float(np.abs(np.diff(grey, axis=0)).mean()), path))
    scored.sort(key=lambda pair: pair[0], reverse=True)
    return [path for _, path in scored[:want]]


def _sku_from_name(name: str) -> str:
    """"BC52-80-TRANG__front.jpg" -> "BC52-80-TRANG". Loose files without one are left."""
    stem = Path(name).stem
    return stem.split("__")[0].strip() if "__" in stem else ""


def write_mapping(report: MediaReport, path: Path) -> None:
    """images.xlsx, so a photo that was renamed can still be traced to its product."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Dữ liệu"
    ws.append(IMAGES.names)
    for sku, names in sorted(report.products.items()):
        for name in names:
            role = name.rsplit("__", 1)[-1].removesuffix(".jpg")
            ws.append([name, sku, role if role in ROLES else "", ""])
    for column, width in zip("ABCD", (34, 24, 12, 40)):
        ws.column_dimensions[column].width = width
    wb.save(path)
