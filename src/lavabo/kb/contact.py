"""Number the photos, so the shop can name them in one message.

Photos arrive from Zalo with names like `received_8837…`, in no order, for products the
sender never mentioned. Asking "what is this one?" per photo is forty messages nobody
answers. Asking once, about a numbered sheet of their own pictures, is one message and one
reply: "3 và 7 là tủ 80 trắng, 4 là gương bo".

So this writes a contact sheet and the mapping table that goes with it. Filling the table
is the only naming step, and it happens once.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

log = logging.getLogger(__name__)

PHOTO_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}
THUMB = 320
COLUMNS = 3
ROWS = 4
LABEL_HEIGHT = 44
MARGIN = 16
MAP_COLUMNS = ["so_thu_tu", "ten_file", "ma_sp", "vai_tro", "ghi_chu"]
MAP_FILENAME = "anh-can-dat-ten.xlsx"


@dataclass(slots=True)
class ContactSheet:
    pages: list[Path]
    mapping: Path
    count: int


def build(source: Path, out: Path) -> ContactSheet:
    from PIL import Image, ImageDraw, ImageFont, ImageOps

    photos = sorted(p for p in source.rglob("*")
                    if p.is_file() and p.suffix.lower() in PHOTO_SUFFIXES)
    if not photos:
        raise ValueError(f"không tìm thấy ảnh nào trong {source}")

    out.mkdir(parents=True, exist_ok=True)
    font = _font(ImageFont)
    per_page = COLUMNS * ROWS
    width = COLUMNS * (THUMB + MARGIN) + MARGIN
    height = ROWS * (THUMB + LABEL_HEIGHT + MARGIN) + MARGIN
    pages: list[Path] = []

    for start in range(0, len(photos), per_page):
        page = Image.new("RGB", (width, height), (255, 255, 255))
        draw = ImageDraw.Draw(page)

        for offset, path in enumerate(photos[start:start + per_page]):
            number = start + offset + 1
            column, row = offset % COLUMNS, offset // COLUMNS
            x = MARGIN + column * (THUMB + MARGIN)
            y = MARGIN + row * (THUMB + LABEL_HEIGHT + MARGIN)

            # A landscape photo does not fill a square cell, and a number drawn at the
            # bottom of the cell then floats nearer the picture BELOW it. On a sheet whose
            # only job is "which number is this photo", that misnames products.
            draw.rectangle([x - 4, y - 4, x + THUMB + 4, y + THUMB + LABEL_HEIGHT],
                           outline=(210, 214, 220), width=2)
            box = (x, y, x + THUMB, y + THUMB)
            try:
                with Image.open(path) as img:
                    thumb = ImageOps.exif_transpose(img).convert("RGB")
                    thumb.thumbnail((THUMB, THUMB))
                    left = x + (THUMB - thumb.width) // 2
                    top = y + (THUMB - thumb.height) // 2
                    page.paste(thumb, (left, top))
                    box = (left, top, left + thumb.width, top + thumb.height)
            except Exception as exc:                   # a broken download in the batch
                log.warning("bỏ qua %s: %s", path.name, exc)
                draw.text((x + 10, y + 10), "?", fill=(150, 150, 150), font=font)

            # Badge on the image itself: unambiguous whatever shape the photo is.
            draw.rectangle([box[0], box[1], box[0] + 54, box[1] + 42], fill=(31, 56, 100))
            draw.text((box[0] + 12, box[1] + 6), f"{number}", fill=(255, 255, 255),
                      font=font)
            caption = _fit(draw, f"{number}. {path.name}", font, THUMB - 12)
            draw.text((x + 6, y + THUMB + 12), caption, fill=(70, 76, 86), font=font)

        target = out / f"anh-danh-so-{len(pages) + 1}.jpg"
        page.save(target, "JPEG", quality=88)
        pages.append(target)

    mapping = _mapping(photos, out / MAP_FILENAME)
    return ContactSheet(pages=pages, mapping=mapping, count=len(photos))


def _fit(draw, text: str, font, width: int) -> str:
    """Trim to the cell. A caption that overflows runs into the neighbouring photo's,
    which on this sheet is the one thing that must never be ambiguous."""
    if draw.textlength(text, font=font) <= width:
        return text
    while text and draw.textlength(text + "…", font=font) > width:
        text = text[:-1]
    return text + "…"


def _font(ImageFont):
    try:
        return ImageFont.load_default(size=30)         # Pillow >= 10.1
    except TypeError:
        return ImageFont.load_default()


def _mapping(photos: list[Path], target: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dữ liệu"
    ws.append(MAP_COLUMNS)
    for number, path in enumerate(photos, start=1):
        ws.append([number, path.name, "", "", ""])
    for column, width in zip("ABCDE", (12, 40, 24, 12, 30)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    wb.save(target)
    return target


def read_mapping(path: Path) -> dict[str, tuple[str, str]]:
    """filename -> (mã SP, vai trò). Rows with no code are simply not yet named."""
    ws = load_workbook(path, data_only=True)["Dữ liệu"]
    header = [str(c).strip().lower() if c is not None else ""
              for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    index = {name: i for i, name in enumerate(header) if name}
    if "ten_file" not in index or "ma_sp" not in index:
        raise ValueError(f"{path.name}: thiếu cột ten_file hoặc ma_sp")

    out: dict[str, tuple[str, str]] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        name = _cell(values, index.get("ten_file"))
        sku = _cell(values, index.get("ma_sp"))
        if name and sku:
            out[name.lower()] = (sku, _cell(values, index.get("vai_tro")))
    return out


def _cell(values, position) -> str:
    if position is None or position >= len(values) or values[position] is None:
        return ""
    return str(values[position]).strip()
