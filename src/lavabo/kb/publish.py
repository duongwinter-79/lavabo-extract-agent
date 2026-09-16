"""A filled intake pack -> the folder Meta's Drive connector reads.

These are different things and conflating them is how a fake product reaches a customer.
The intake pack is a FORM: placeholder markers, grey example rows, a VI-DU product, and
tabs that exist for us rather than for the shop's customers. The connected folder is
PUBLISHED CONTENT: every line in it is something the agent may say out loud.

So publishing is a filter, not a copy. It drops:

  voice.md, dont_say.md   instructions, not knowledge -- they belong in Hướng dẫn, and an
                          agent that reads them as facts will recite house rules at a
                          customer (docs/18 §1)
  synonyms.xlsx           retrieval hints for our own search; meaningless as knowledge
  images.xlsx             a filename mapping. The agent would read "IMG_4821.jpg" as a fact
  example rows            anything still carrying the template's VI-DU marker

and it refuses to run at all while `kb check` has a fatal problem, for the same reason
`kb feed` does: this is the moment a spreadsheet becomes something the Page says.
"""

from __future__ import annotations

import logging
import re
import shutil
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .spec import CATALOG, EXAMPLE_MARKER, FAQ, PLACEHOLDER, PROMOTIONS, SHIPPING

log = logging.getLogger(__name__)

KNOWLEDGE_DIR = "01-KIEN-THUC-AI"
IMAGES_DIR = "02-ANH-SAN-PHAM"
FORM_DIR = "bieu-mau"

# Columns a customer may be told about. `anh` is filenames and `cap_nhat_ngay` is our
# bookkeeping -- both read as noise, or worse as facts, to an agent.
PRICE_COLUMNS = ["ma_sp", "ten_sp", "loai", "kich_thuoc", "chat_lieu", "mau", "don_vi",
                 "gia_niem_yet", "gia_km", "km_den_ngay", "tinh_trang", "thoi_gian_giao",
                 "bao_hanh", "gia_gom", "ghi_chu_tu_van"]

# Only ever written when --image-base says the photos are reachable; see `photo_urls`.
PHOTO_COLUMN = "link_anh"

EXCLUDED = {
    "voice.md": "hướng dẫn cách trả lời — dán vào tab Hướng dẫn, không phải kiến thức",
    "dont_say.md": "hướng dẫn — dán vào tab Hướng dẫn",
    "handoff.xlsx": "danh sách chuyển cho người thật — dán vào tab Hướng dẫn",
    "synonyms.xlsx": "chỉ dùng cho tìm kiếm nội bộ",
    "images.xlsx": "bảng ánh xạ tên file ảnh, không phải thông tin cho khách",
}

# Goes in the PARENT folder, never the connected one: a note about folder rules, indexed
# as knowledge, becomes the agent explaining our filing system to a customer.
PARENT_README = """THƯ MỤC NÀY LÀ GÌ

01-KIEN-THUC-AI   ĐÃ NỐI với Meta Business Agent. AI đọc TẤT CẢ file trong
                  đây và có thể nói lại với khách. Chỉ để file đã kiểm tra.
02-ANH-SAN-PHAM   Ảnh sản phẩm. KHÔNG nối với AI.
bieu-mau          Biểu mẫu để điền. KHÔNG nối với AI — trong đó còn chỗ
                  trống và dòng ví dụ mẫu.

QUY TẮC
- Sửa thẳng vào file cũ. Không tạo "bang-gia-v2", "moi-nhat".
- Mỗi file có dòng "Cập nhật ngày ..." ở đầu — sửa dòng đó mỗi lần kiểm tra lại.
- Giá chỉ nằm trong 05-bang-gia. Không ghi giá ở file khác.
"""


# What publishing actually writes. A fatal problem in one of these blocks it; a problem
# in a file that is never published (voice.md, dont_say.md) does not -- it blocks turning
# the agent ON, which is a different action with a different blast radius.
PUBLISHED_SOURCES = {"store.md", "policies.md", SHIPPING.filename, FAQ.filename,
                     CATALOG.filename, PROMOTIONS.filename}


def blocking(problems: list) -> list:
    """The fatal problems that stand between here and a published folder."""
    return [p for p in problems if p.fatal and p.file in PUBLISHED_SOURCES]


def instruction_gaps(problems: list) -> list:
    """Unanswered questions in the files that become Hướng dẫn. Not published, still
    required before the agent may be switched on."""
    return [p for p in problems if p.fatal and p.file in ("voice.md", "dont_say.md")]


@dataclass
class PublishReport:
    root: Path
    written: list[str] = field(default_factory=list)
    excluded: list[str] = field(default_factory=list)
    skipped_examples: int = 0
    images: int = 0


def publish(intake: Path, out: Path, *, today: date | None = None,
            image_base: str = "") -> PublishReport:
    stamp = (today or date.today()).isoformat()
    knowledge = out / KNOWLEDGE_DIR
    knowledge.mkdir(parents=True, exist_ok=True)
    (out / IMAGES_DIR).mkdir(exist_ok=True)
    (out / FORM_DIR).mkdir(exist_ok=True)
    (out / "00-DOC-TRUOC.txt").write_text(PARENT_README, encoding="utf-8")

    report = PublishReport(root=out)

    if (intake / "store.md").exists():
        _prose(intake / "store.md", knowledge / "01-thong-tin-cua-hang.md",
               "Thông tin cửa hàng", stamp, report)
    if (intake / "policies.md").exists():
        _prose(intake / "policies.md", knowledge / "02-chinh-sach.md",
               "Chính sách", stamp, report)

    _table(intake / SHIPPING.filename, knowledge / "03-phi-van-chuyen.xlsx",
           SHIPPING.names, stamp, report)
    _table(intake / FAQ.filename, knowledge / "04-cau-hoi-thuong-gap.xlsx",
           FAQ.names, stamp, report)
    _table(intake / CATALOG.filename, knowledge / "05-bang-gia.xlsx",
           PRICE_COLUMNS, stamp, report,
           derived=(PHOTO_COLUMN, lambda rec: photo_urls(image_base, rec.get("anh")))
           if image_base else None)
    _table(intake / PROMOTIONS.filename, knowledge / "06-khuyen-mai.xlsx",
           PROMOTIONS.names, stamp, report, optional=True)

    report.images = _images(intake / "images", out / IMAGES_DIR)
    report.excluded = [f"{name} — {why}" for name, why in EXCLUDED.items()
                       if (intake / name).exists()]
    return report


# ------------------------------------------------------------------------------ prose

def _prose(source: Path, target: Path, title: str, stamp: str,
           report: PublishReport) -> None:
    """Strip the form furniture, keep the answers.

    A filled `- **Hotline:** 0912 345 678   ← ví dụ: ...` becomes `- Hotline: 0912 345 678`.
    The example hint has to go: left in, it is a second phone number in the knowledge.
    """
    out = [f"# {title}", "", f"Cập nhật ngày: {stamp}", ""]

    for line in source.read_text(encoding="utf-8").splitlines():
        text = line.rstrip()
        if text.startswith(">") or text.startswith("# "):
            continue                                   # instructions to the shop
        if PLACEHOLDER in text:
            continue                                   # never answered
        if text.strip().startswith("*") and text.strip().endswith("*"):
            continue                                   # our italic notes
        if text.startswith("## "):
            out += ["", f"## {text[3:].strip()}", ""]
            continue

        text = re.sub(r"\s*←.*$", "", text)            # "← ví dụ: ..."
        text = text.replace("*(không bắt buộc)*", "")
        text = text.replace("**", "")
        if text.strip():
            out.append(text.rstrip())

    target.write_text("\n".join(out).rstrip() + "\n", encoding="utf-8")
    report.written.append(target.name)
    log.info("wrote %s", target)


# ----------------------------------------------------------------------------- tables

def photo_urls(base: str, names) -> str:
    """The photos' public addresses, from the filenames the catalogue already carries.

    A bare filename is excluded from knowledge on purpose -- `images.xlsx` is dropped
    because the agent would state "IMG_4821.jpg" as though it meant something. A URL is the
    same datum made actionable: a customer can open it. So it earns a place in the
    knowledge only once `--image-base` asserts the photos are actually reachable, and never
    by default.
    """
    parts = [n.strip() for n in str(names or "").split(";") if n.strip()]
    return "; ".join(f"{base.rstrip('/')}/{name.lstrip('/')}" for name in parts)


def _table(source: Path, target: Path, columns: list[str], stamp: str,
           report: PublishReport, *, optional: bool = False,
           derived: tuple[str, Callable[[dict], object]] | None = None) -> None:
    """`derived` appends one column computed from the source row, for a value the shop
    never types -- the photo's public URL, which does not exist until somebody hosts it."""
    if not source.exists():
        if not optional:
            log.warning("thiếu %s", source.name)
        return

    ws_in = load_workbook(source, data_only=True)
    sheet = ws_in["Dữ liệu"] if "Dữ liệu" in ws_in.sheetnames else ws_in.worksheets[0]
    header = [str(c).strip().lower() if c is not None else ""
              for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    index = {name: i for i, name in enumerate(header) if name}

    wb = Workbook()
    out = wb.active
    out.title = "Dữ liệu"
    out.append([f"Cập nhật ngày: {stamp}"])
    out.append([c for c in columns if c in index] + ([derived[0]] if derived else []))

    rows = 0
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        record = {name: values[at] if at < len(values) else None
                  for name, at in index.items()}
        row = [record[c] for c in columns if c in index]
        if any(EXAMPLE_MARKER in str(v) for v in row if v is not None):
            report.skipped_examples += 1
            continue
        if derived:
            row.append(derived[1](record))
        out.append(row)
        rows += 1

    if not rows and optional:
        return

    for idx in range(1, len(out[2]) + 1):
        out.column_dimensions[get_column_letter(idx)].width = 22
    wb.save(target)
    report.written.append(target.name)
    log.info("wrote %s (%d rows)", target, rows)


def _images(source: Path, target: Path) -> int:
    if not source.is_dir():
        return 0
    count = 0
    for path in sorted(source.iterdir()):
        if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}:
            shutil.copy2(path, target / path.name)
            count += 1
    return count
