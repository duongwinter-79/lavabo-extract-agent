"""The pack as one workbook, and the rules following it there.

The point of this format is a shop owner with a phone and no Excel: one Google Sheets
link, filled in a browser. The risk is having two packs whose rules quietly diverge, so
these tests are mostly about the single workbook reaching the SAME verdict as the folder.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook                                # noqa: E402

from lavabo.kb import spec                                        # noqa: E402
from lavabo.kb.check import check_intake, check_one_file          # noqa: E402
from lavabo.kb.onefile import (ANSWER_COLUMN, DOC_TABS, TABS,  # noqa: E402
                               write_one_file, write_pack)

from test_kb_check import GOOD                                    # noqa: E402


class TheWorkbook(unittest.TestCase):
    def setUp(self):
        self.path = Path(tempfile.mkdtemp()) / "pack.xlsx"
        write_one_file(self.path)

    def test_every_form_is_a_tab(self):
        names = load_workbook(self.path).sheetnames
        self.assertEqual(names[0], "Đọc trước")
        for title in list(DOC_TABS) + list(TABS):
            self.assertIn(title, names)

    def test_the_data_tabs_carry_the_same_columns_as_the_folder_pack(self):
        book = load_workbook(self.path)
        for title, sheet in TABS.items():
            with self.subTest(title):
                header = [c.value for c in book[title][1]]
                self.assertEqual(header, sheet.names)

    def test_the_handoff_tab_carries_its_seeded_rows(self):
        """The one-file build had its own copy of the row-seeding logic, so this tab
        shipped with a single example row where twelve rules should be."""
        ws = load_workbook(self.path)["Không được tự trả lời"]
        topics = [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        self.assertEqual(len(topics), len(spec.LOCKED_TOPICS))
        self.assertIn("Khiếu nại, hàng lỗi, hàng vỡ", topics)

    def test_seeded_tabs_carry_no_example_row(self):
        ws = load_workbook(self.path)["Không được tự trả lời"]
        cells = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c]
        self.assertFalse(any(spec.EXAMPLE_MARKER in c for c in cells))

    def test_it_refuses_to_overwrite_without_force(self):
        with self.assertRaises(FileExistsError):
            write_one_file(self.path)
        write_one_file(self.path, force=True)

    def test_an_untouched_workbook_names_the_missing_answers(self):
        fatal = [p for p in check_one_file(self.path) if p.fatal]
        self.assertEqual(len(fatal), len(DOC_TABS))
        self.assertTrue(any("Tên shop" in str(p) for p in fatal))

    def test_problems_are_located_by_tab_not_by_filename(self):
        problems = [str(p) for p in check_one_file(self.path)]
        self.assertTrue(any("tab 'Cửa hàng'" in p for p in problems), problems)
        self.assertFalse(any("store.md" in p for p in problems), problems)


class FillingItIn(unittest.TestCase):
    def setUp(self):
        self.path = Path(tempfile.mkdtemp()) / "pack.xlsx"
        write_one_file(self.path)
        self.book = load_workbook(self.path)

    def save(self):
        self.book.save(self.path)

    def answer_everything(self):
        for title, doc in DOC_TABS.items():
            ws = self.book[title]
            for row in ws.iter_rows(min_row=2):
                label = row[0].value
                if label and not str(label).startswith(("—", doc.why[:20])):
                    ws.cell(row=row[0].row, column=ANSWER_COLUMN, value="xong")

    def write_catalog(self, row: dict):
        ws = self.book["Danh mục sản phẩm"]
        ws.delete_rows(2, ws.max_row)
        ws.append([row.get(c) for c in spec.CATALOG.names])

    def test_a_fully_answered_workbook_has_no_fatal_problems(self):
        self.answer_everything()
        self.write_catalog(GOOD)
        self.save()
        fatal = [str(p) for p in check_one_file(self.path) if p.fatal]
        self.assertEqual(fatal, [])

    def test_a_bad_price_is_caught_in_the_workbook_too(self):
        """Google Sheets drops the whole-number guard the Excel form has, so this is the
        only thing standing between "2tr850" and a feed."""
        self.answer_everything()
        self.write_catalog({**GOOD, "gia_niem_yet": "2tr850"})
        self.save()
        fatal = [str(p) for p in check_one_file(self.path) if p.fatal]
        self.assertEqual(len(fatal), 1)
        self.assertIn("2850000", fatal[0])
        self.assertIn("Danh mục sản phẩm", fatal[0])

    def test_an_expired_promo_is_caught_in_the_workbook_too(self):
        self.answer_everything()
        self.write_catalog({**GOOD, "gia_km": 2550000,
                            "km_den_ngay": date(2020, 1, 1)})
        self.save()
        self.assertTrue(any("hết hạn" in str(p)
                            for p in check_one_file(self.path) if p.fatal))

    def test_a_placeholder_typed_back_in_still_counts_as_empty(self):
        """Somebody will paste the marker into the answer column. It is not an answer."""
        self.answer_everything()
        ws = self.book["Cửa hàng"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "Tên shop":
                ws.cell(row=row[0].row, column=ANSWER_COLUMN, value=spec.PLACEHOLDER)
        self.write_catalog(GOOD)
        self.save()
        fatal = [str(p) for p in check_one_file(self.path) if p.fatal]
        self.assertTrue(any("Tên shop" in f for f in fatal), fatal)


class SpreadingItBackOut(unittest.TestCase):
    """`write_pack`: the workbook back into the folder everything downstream reads.

    `check` takes either shape, but publish and feed only take the folder, so a shop that
    filled the workbook could reach a passing check and go no further. These tests are
    about the bridge carrying the answers across without changing the verdict.
    """

    def setUp(self):
        self.path = Path(tempfile.mkdtemp()) / "pack.xlsx"
        write_one_file(self.path)
        book = load_workbook(self.path)
        for title, doc in DOC_TABS.items():
            ws = book[title]
            for row in ws.iter_rows(min_row=2):
                label = row[0].value
                if label and not str(label).startswith(("—", doc.why[:20])):
                    ws.cell(row=row[0].row, column=ANSWER_COLUMN, value="xong")
        catalog = book["Danh mục sản phẩm"]
        catalog.delete_rows(2, catalog.max_row)
        catalog.append([GOOD.get(c) for c in spec.CATALOG.names])
        book.save(self.path)

    def pack(self, images=None):
        return write_pack(self.path, Path(tempfile.mkdtemp()) / "pack", images=images)

    def test_the_folder_it_writes_passes_the_check_the_workbook_passed(self):
        """One source, one verdict -- the whole reason this function exists."""
        self.assertEqual([str(p) for p in check_intake(self.pack()) if p.fatal], [])

    def test_every_file_the_folder_pack_expects_is_written(self):
        pack = self.pack()
        for sheet in TABS.values():
            self.assertTrue((pack / sheet.filename).is_file(), sheet.filename)
        for doc in DOC_TABS.values():
            self.assertTrue((pack / doc.filename).is_file(), doc.filename)

    def test_an_answer_survives_the_trip_into_the_markdown(self):
        self.assertIn("xong", (self.pack() / "store.md").read_text(encoding="utf-8"))

    def test_an_unanswered_field_keeps_its_placeholder_for_check_to_find(self):
        """Filling a blank with something plausible would hide the gap, not close it."""
        book = load_workbook(self.path)
        ws = book["Cửa hàng"]
        for row in ws.iter_rows(min_row=2):
            if str(row[0].value or "").strip() == "Hotline":
                ws.cell(row=row[0].row, column=ANSWER_COLUMN, value=None)
        book.save(self.path)
        text = (self.pack() / "store.md").read_text(encoding="utf-8")
        self.assertIn(spec.PLACEHOLDER, text)

    def test_photos_are_copied_only_when_a_folder_is_given(self):
        """The workbook carries no images, so without --images there are none to publish."""
        self.assertEqual(list((self.pack() / "images").iterdir()), [])

        photos = Path(tempfile.mkdtemp())
        (photos / "tu-80.jpg").write_bytes(b"x")
        (photos / "ghi-chu.txt").write_bytes(b"x")          # not a photograph
        copied = [p.name for p in (self.pack(images=photos) / "images").iterdir()]
        self.assertEqual(copied, ["tu-80.jpg"])


class ItIsMissing(unittest.TestCase):
    def test_a_path_that_does_not_exist_is_reported_not_raised(self):
        problems = check_one_file(Path(tempfile.mkdtemp()) / "nope.xlsx")
        self.assertEqual(len(problems), 1)
        self.assertTrue(problems[0].fatal)


if __name__ == "__main__":
    unittest.main()
