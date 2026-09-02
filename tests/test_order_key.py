"""Which orders are the same order.

An order was identified by (ngày, tháng, số đơn), which holds only while ONE person posts
orders. Copy-paste carries no sender, so that assumption was never chosen -- it was
invisible. Two people each numbering their own orders from 1 collide on the first day of
the month, and the second order merged into the first or was filed as a competing version
of it.

The key now carries who the order belongs to: the reporter when the chat names one, the
customer when it does not.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

import zalo_capture as zc                                         # noqa: E402
from lavabo import reporters                                      # noqa: E402
from lavabo.config import Config                                  # noqa: E402

ORDER = "1 tủ BC52\nThái Thuỵ, TB\nTổng 5.800\nĐã cọc 500k"


class ReadingTheReporter(unittest.TestCase):
    def test_a_name_before_the_header_is_the_sender(self):
        for line in ("Trà My: 13/7 đơn 1 - Chị Hương",
                     "Ngọc Anh (11:52): 13/7 đơn 1 - Chị Hương"):
            with self.subTest(line):
                who, rest = zc.split_reporter(line)
                self.assertIn(who, ("Trà My", "Ngọc Anh"))
                self.assertTrue(rest.strip().startswith("13/7"))

    def test_a_customer_after_a_colon_is_not_a_sender(self):
        """"17/7 đơn 1: Thảo Nguyên" puts the CUSTOMER after a colon. Reading that as a
        sender would rename every order written that way."""
        who, rest = zc.split_reporter("17/7 đơn 1: Thảo Nguyên")
        self.assertIsNone(who)
        self.assertEqual(rest, "17/7 đơn 1: Thảo Nguyên")

    def test_a_prefix_that_reveals_no_order_is_left_alone(self):
        who, _ = zc.split_reporter("Trà My: ok chị em nhận rồi")
        self.assertIsNone(who)

    def test_a_standalone_name_needs_a_marker(self):
        """Ordinary chatter sits right before the next order all the time. A wrong
        reporter goes into the key and splits an order that should have merged, so this
        errs towards missing one -- a missed reporter still falls back to the customer."""
        self.assertTrue(zc.REPORTER_LINE.match("Trà My:"))
        self.assertTrue(zc.REPORTER_LINE.match("Ngọc Anh (11:52)"))
        self.assertIsNone(zc.REPORTER_LINE.match("ok chị em nhận rồi"))
        self.assertIsNone(zc.REPORTER_LINE.match("vâng ạ"))

    def test_a_sender_between_orders_is_not_swallowed_by_the_one_above(self):
        blocks = zc.split_orders(
            f"Trà My: 13/7 đơn 1 - Chị Hương\n{ORDER}\n"
            f"Ngọc Anh (11:52)\n13/7 đơn 1 - Anh Lợi\n{ORDER}", target_month=7)
        self.assertEqual([b.reporter for b in blocks], ["Trà My", "Ngọc Anh"])
        self.assertNotIn("Ngọc Anh (11:52)", blocks[0].lines)


class ReporterInTheHeader(unittest.TestCase):
    """The format the shop is adopting, because copy-paste carries no sender: the name is
    typed into the header itself, last, after the customer."""

    def _parse(self, line, known=None):
        from lavabo.connectors.zalo_export import ORDER_HEADER, header_parties

        match = ORDER_HEADER.match(line)
        self.assertIsNotNone(match, f"header did not match: {line!r}")
        return header_parties(match, known)

    def test_the_dashed_form_splits_into_customer_and_reporter(self):
        self.assertEqual(self._parse("13/7 đơn 1 - Chị Hương - Trà My"),
                         ("Chị Hương", "Trà My"))

    def test_the_bracketed_form_does_too(self):
        """This one matched NOTHING before -- ")" ended the group, so the whole order
        was dropped rather than mis-parsed."""
        self.assertEqual(self._parse("13/7 đơn 2 (Trần Thị Liên) - Ngọc Anh"),
                         ("Trần Thị Liên", "Ngọc Anh"))

    def test_headers_the_shop_already_writes_are_unchanged(self):
        for line, customer in (("13/7 đơn 1 - Chị Hương", "Chị Hương"),
                               ("2/7 đơn 2 (Trần Thị Liên)", "Trần Thị Liên"),
                               ("29/6 đơn 4- Bùi Đức Hạnh", "Bùi Đức Hạnh"),
                               ("16/7 đơn 1 - Tiểu bảo bối", "Tiểu bảo bối"),
                               ("17/7 đơn 1: Thảo Nguyên", "Thảo Nguyên")):
            with self.subTest(line):
                self.assertEqual(self._parse(line), (customer, None))

    def test_a_ruling_settles_an_ambiguous_trailing_field(self):
        """"Anh Tâm - Hà Nội" is one customer and a place; "Anh Tâm - Trà My" is a
        customer and a staff member. Nothing in the text tells them apart -- only a
        verdict about the NAME does, and both verdicts have to be storable."""
        decided = {"tra my": True, "ha noi": False}
        self.assertEqual(self._parse("13/7 đơn 3 - Anh Tâm - Hà Nội", decided),
                         ("Anh Tâm - Hà Nội", None))
        self.assertEqual(self._parse("13/7 đơn 3 - Anh Tâm - Trà My", decided),
                         ("Anh Tâm", "Trà My"))

    def test_an_unruled_name_still_reads_as_a_reporter(self):
        """A new staff member's first order must not have their name folded into the
        customer just because OTHER names are already known."""
        self.assertEqual(self._parse("13/7 đơn 3 - Anh Tâm - Bảo Ngọc", {"tra my": True}),
                         ("Anh Tâm", "Bảo Ngọc"))

    def test_the_header_beats_a_sender_line(self):
        blocks = zc.split_orders(
            f"Ngọc Anh: 13/7 đơn 1 - Chị Hương - Trà My\n{ORDER}", target_month=7)
        self.assertEqual(blocks[0].reporter, "Trà My")
        self.assertEqual(blocks[0].customer, "Chị Hương")


class Identity(unittest.TestCase):
    def _block(self, reporter=None, customer="Chị Hương", order_no=1):
        header = f"13/7 đơn {order_no}" + (f" - {customer}" if customer else "")
        text = (f"{reporter}: {header}\n{ORDER}" if reporter else f"{header}\n{ORDER}")
        return zc.split_orders(text, target_month=7)[0]

    def test_the_reporter_decides_when_the_chat_names_one(self):
        self.assertEqual(self._block(reporter="Trà My").key, (13, 7, 1, "tra my"))

    def test_the_customer_stands_in_when_it_does_not(self):
        self.assertEqual(self._block().key, (13, 7, 1, "chi huong"))

    def test_names_are_folded_so_one_order_stays_one_order(self):
        self.assertEqual(self._block(customer="chi Huong").key,
                         self._block(customer="Chị Hương").key)

    def test_two_people_on_one_order_number_are_two_orders(self):
        a = self._block(reporter="Trà My", customer="Chị Hương")
        b = self._block(reporter="Ngọc Anh", customer="Anh Lợi")
        self.assertNotEqual(a.key, b.key)

    def test_two_customers_on_one_order_number_are_two_orders(self):
        """Even with no sender in the chat at all -- which is every capture made so far."""
        self.assertNotEqual(self._block(customer="Chị Hương").key,
                            self._block(customer="Anh Lợi").key)


class Capturing(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cfg = Config()
        self.cfg.zalo.inbox_dir = Path(self.tmp.name) / "zalo"
        self.cfg.zalo.inbox_dir.mkdir(parents=True)

    def tearDown(self):
        self.tmp.cleanup()

    def _capture(self, text):
        return zc.handle_orders(text, self.cfg, 7, 2026, all_months=False, trim=True,
                                closer="Trà My")

    def _files(self):
        return sorted(p.name for p in self.cfg.zalo.inbox_dir.glob("*.txt"))

    def test_two_people_both_posting_don_1_produce_two_orders(self):
        result = self._capture(
            f"Trà My: 13/7 đơn 1 - Chị Hương\n{ORDER}\n"
            f"Ngọc Anh (11:52)\n13/7 đơn 1 - Anh Lợi\n{ORDER}")
        self.assertEqual(result.saved, 2)
        self.assertEqual(result.versions, 0, "these are two orders, not two versions")
        self.assertEqual(len(self._files()), 2)

    def test_re_pasting_that_chat_adds_nothing(self):
        chat = (f"Trà My: 13/7 đơn 1 - Chị Hương\n{ORDER}\n"
                f"Ngọc Anh (11:52)\n13/7 đơn 1 - Anh Lợi\n{ORDER}")
        self._capture(chat)
        again = self._capture(chat)
        self.assertEqual((again.saved, again.duplicates), (0, 2))

    def test_the_reporter_is_remembered_beside_the_order(self):
        self._capture(f"Trà My: 13/7 đơn 1 - Chị Hương\n{ORDER}")
        self.assertEqual(list(reporters.load(self.cfg.zalo.inbox_dir).values()), ["Trà My"])

    def test_an_order_captured_before_senders_existed_is_not_duplicated_by_one(self):
        """The migration case. Every order already captured is filed under its customer;
        the first paste after the chat starts naming senders brings a reporter with it.
        Matching on either name is what stops that re-filing the whole month."""
        self._capture(f"13/7 đơn 1 - Chị Hương\n{ORDER}")
        after = self._capture(f"Trà My: 13/7 đơn 1 - Chị Hương\n{ORDER}")
        self.assertEqual((after.saved, after.duplicates), (0, 1))
        self.assertEqual(len(self._files()), 1)

    def test_and_it_learns_the_reporter_from_that_paste(self):
        self._capture(f"13/7 đơn 1 - Chị Hương\n{ORDER}")
        self._capture(f"Trà My: 13/7 đơn 1 - Chị Hương\n{ORDER}")
        self.assertEqual(list(reporters.load(self.cfg.zalo.inbox_dir).values()), ["Trà My"])

    def test_the_header_format_separates_two_people(self):
        result = self._capture(f"13/7 đơn 1 - Chị Hương - Trà My\n{ORDER}\n"
                               f"13/7 đơn 1 - Anh Lợi - Ngọc Anh\n{ORDER}")
        self.assertEqual(result.saved, 2)
        self.assertEqual(len(self._files()), 2)

    def test_the_customer_column_does_not_swallow_the_reporter(self):
        self._capture(f"13/7 đơn 1 - Chị Hương - Trà My\n{ORDER}")
        block = zc.split_orders(f"13/7 đơn 1 - Chị Hương - Trà My\n{ORDER}",
                                target_month=7, inbox=self.cfg.zalo.inbox_dir)[0]
        self.assertEqual(block.customer, "Chị Hương")
        self.assertEqual(block.reporter, "Trà My")

    def test_a_reporter_is_never_erased_by_a_copy_that_lost_the_names(self):
        self._capture(f"Trà My: 13/7 đơn 1 - Chị Hương\n{ORDER}")
        self._capture(f"13/7 đơn 1 - Chị Hương\n{ORDER}")
        self.assertEqual(list(reporters.load(self.cfg.zalo.inbox_dir).values()), ["Trà My"])


class SettlingAmbiguousNames(unittest.TestCase):
    """No rule separates "Anh Tâm - Hà Nội" from "Anh Tâm - Trà My": both are two short
    names either side of a dash, and "Trà My" is a district in Quảng Nam as well as a
    person here. The question is about the world, so it is asked once per NAME and the
    answer is kept -- which makes the deterministic reading exact instead of replacing it.
    """

    BODY = "1 tủ BC52\nTổng 5.800\nĐã cọc 500k"

    def setUp(self):
        from lavabo import segment

        self.tmp = tempfile.TemporaryDirectory()
        self.cfg = Config()
        self.cfg.zalo.inbox_dir = Path(self.tmp.name) / "zalo"
        self.cfg.zalo.inbox_dir.mkdir(parents=True)
        self.cfg.extract.ai_segmentation = "shadow"
        self.calls = 0
        self.segment = segment
        self._real = segment.completer_for
        segment.completer_for = lambda cfg: self._fake()

    def tearDown(self):
        self.segment.completer_for = self._real
        self.tmp.cleanup()

    def _fake(self, verdicts=(("Hà Nội", "customer"), ("Bảo Ngọc", "staff"))):
        outer = self

        class Fake:
            def complete_json(self, system, user, schema, *, max_tokens=0):
                if "names" in schema.get("properties", {}):
                    outer.calls += 1
                    return outer.segment.Completion(
                        {"names": [{"name": n, "role": r} for n, r in verdicts]},
                        300, 40, "STOP")
                return outer.segment.Completion({"orders": []}, 100, 20, "STOP")
        return Fake()

    def _chat(self):
        return (f"13/7 đơn 3 - Anh Tâm - Hà Nội\n{self.BODY}\n"
                f"13/7 đơn 1 - Chị Hương - Bảo Ngọc\n{self.BODY}\n"
                f"14/7 đơn 1 - Minh Nguyễn - Bảo Ngọc\n{self.BODY}")

    def _capture(self, text=None):
        return zc.handle_orders(text or self._chat(), self.cfg, 7, 2026,
                                all_months=False, trim=True, closer="Trà My")

    def test_only_genuinely_ambiguous_names_are_asked_about(self):
        known = {"tra my": True}
        candidates = zc.undecided_names(
            "13/7 đơn 1 - Chị Hương - Trà My\n"
            "13/7 đơn 3 - Anh Tâm - Hà Nội\n"
            "13/7 đơn 4 - Chị Lan\n"
            "13/7 đơn 5 - Anh Tâm - 0912345678", known)
        self.assertEqual(set(candidates), {"Hà Nội"})

    def test_a_place_is_left_with_the_customer(self):
        self._capture()
        block = zc.split_orders(f"13/7 đơn 3 - Anh Tâm - Hà Nội\n{self.BODY}",
                                target_month=7, inbox=self.cfg.zalo.inbox_dir)[0]
        self.assertEqual(block.customer, "Anh Tâm - Hà Nội")
        self.assertIsNone(block.reporter)

    def test_a_staff_name_becomes_the_reporter(self):
        self._capture()
        block = zc.split_orders(f"13/7 đơn 1 - Chị Hương - Bảo Ngọc\n{self.BODY}",
                                target_month=7, inbox=self.cfg.zalo.inbox_dir)[0]
        self.assertEqual((block.customer, block.reporter), ("Chị Hương", "Bảo Ngọc"))

    def test_the_question_is_asked_once_and_then_never_again(self):
        self._capture()
        self.assertEqual(self.calls, 1)
        self._capture()
        self._capture()
        self.assertEqual(self.calls, 1, "a settled name must not be re-asked")

    def test_both_verdicts_are_stored(self):
        """Knowing a name is NOT staff is worth as much as knowing it is: without it the
        shape test splits "Anh Tâm - Hà Nội" again on every paste."""
        from lavabo import staff

        self._capture()
        stored = staff.load(self.cfg.zalo.inbox_dir)
        self.assertEqual(stored.get("Hà Nội"), staff.NOT_STAFF)
        self.assertEqual(stored.get("Bảo Ngọc"), staff.AI)

    def test_a_person_outranks_the_model(self):
        from lavabo import staff

        self._capture()
        staff.record(self.cfg.zalo.inbox_dir, "Hà Nội", staff.TYPED)
        self.assertEqual(staff.load(self.cfg.zalo.inbox_dir)["Hà Nội"], staff.TYPED)

    def test_an_unreachable_model_leaves_the_plain_reading_alone(self):
        """Failing to reach the model must keep the existing guess, not invent another."""
        self.segment.completer_for = lambda cfg: None
        self._capture()
        block = zc.split_orders(f"13/7 đơn 3 - Anh Tâm - Hà Nội\n{self.BODY}",
                                target_month=7, inbox=self.cfg.zalo.inbox_dir)[0]
        self.assertEqual((block.customer, block.reporter), ("Anh Tâm", "Hà Nội"))

    def test_nothing_is_asked_when_no_name_is_ambiguous(self):
        self._capture(f"13/7 đơn 1 - Chị Hương\n{self.BODY}")
        self.assertEqual(self.calls, 0)


class GapWarning(unittest.TestCase):
    """order_gaps unpacked the order key positionally, so widening the key crashed the
    status endpoint the page polls -- the whole app, on every refresh, for anyone with a
    single order captured. No test touched it, which is why it shipped."""

    BODY = "1 tủ\nTổng 5.800"

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cfg = Config()
        self.cfg.zalo.inbox_dir = Path(self.tmp.name) / "zalo"
        self.cfg.zalo.inbox_dir.mkdir(parents=True)

    def tearDown(self):
        self.tmp.cleanup()

    def _capture(self, header):
        zc.handle_orders(f"{header}\n{self.BODY}", self.cfg, 8, 2026,
                         all_months=False, trim=True, closer="Trà My")

    def _gaps(self):
        return zc.order_gaps(self.cfg.zalo.inbox_dir, 8)

    def test_it_runs_at_all(self):
        """The crash: too many values to unpack."""
        self._capture("15/8 đơn 1 - Chị Hương")
        self.assertEqual(self._gaps(), [])

    def test_a_missing_number_is_reported(self):
        self._capture("15/8 đơn 2 - KH2")
        self._capture("15/8 đơn 3 - KH3")
        self.assertEqual(len(self._gaps()), 1)
        self.assertIn("thiếu đơn 1", self._gaps()[0])

    def test_two_people_each_complete_raise_no_alarm(self):
        self._capture("15/8 đơn 1 - KH a - Trà My")
        self._capture("15/8 đơn 2 - KH b - Trà My")
        self._capture("15/8 đơn 1 - KH c - Ngọc Anh")
        self.assertEqual(self._gaps(), [])

    def test_one_persons_gap_is_not_filled_in_by_the_other(self):
        """Trà My's 1 and 3 plus Ngọc Anh's 2 reads as an unbroken 1, 2, 3 when the day
        is counted as one sequence, and the missing order is never reported."""
        self._capture("15/8 đơn 1 - KH a - Trà My")
        self._capture("15/8 đơn 3 - KH b - Trà My")
        self._capture("15/8 đơn 2 - KH c - Ngọc Anh")
        gaps = " | ".join(self._gaps())
        self.assertIn("Trà My", gaps)
        self.assertIn("thiếu đơn 2", gaps)

    def test_another_month_is_not_counted(self):
        self._capture("15/8 đơn 2 - KH2")
        self.assertEqual(zc.order_gaps(self.cfg.zalo.inbox_dir, 7), [])


class Flagging(unittest.TestCase):
    """Both rows are written either way; the sheet has to say which kind of clash it is."""

    def _sheet(self, orders):
        import tempfile as tf

        from openpyxl import load_workbook

        from lavabo.load.senkahomes import COL_REVIEW, write_orders_workbook
        from lavabo.models import Conversation, ExtractionResult, Source

        convs, results = [], {}
        for cid, day, no, customer, reporter in orders:
            conv = Conversation(conversation_id=cid, source=Source.ZALO,
                                customer_name=customer, messages=[])
            conv.raw = {"order_day": day, "order_month": 7, "order_number": no}
            if reporter:
                conv.raw["order_reporter"] = reporter
            convs.append(conv)
            results[cid] = ExtractionResult(
                conversation_id=cid, source=Source.ZALO,
                values={"items": [{"name": "tủ", "quantity": 1}], "address": "HN",
                        "total_text": "5tr", "deposit_text": "500k"})

        with tf.TemporaryDirectory() as tmp:
            out = Path(tmp) / "o.xlsx"
            write_orders_workbook(out, convs, results, default_year=2026)
            ws = load_workbook(out).active
            return [(r[2].value, r[COL_REVIEW - 1].value, r[2].fill.fgColor.rgb)
                    for r in ws.iter_rows(min_row=2) if r[0].value is not None]

    def test_different_people_are_marked_as_a_numbering_clash(self):
        rows = self._sheet([("a", 13, 1, "Chị Hương", "Trà My"),
                            ("b", 13, 1, "Anh Lợi", "Ngọc Anh")])
        self.assertEqual(len(rows), 2, "both orders are real and both get a row")
        for _, reason, fill in rows:
            self.assertIn("khác người/khách", reason)
            self.assertNotEqual(fill, "00000000", "and both are highlighted")

    def test_the_same_order_twice_is_marked_as_a_duplicate(self):
        rows = self._sheet([("a", 14, 2, "Minh", "Trà My"),
                            ("b", 14, 2, "Minh", "Trà My")])
        for _, reason, _ in rows:
            self.assertEqual(reason, "trùng số đơn")

    def test_an_unambiguous_order_is_not_flagged(self):
        rows = self._sheet([("a", 15, 3, "Hoà", "Trà My")])
        self.assertIsNone(rows[0][1])
        self.assertEqual(rows[0][2], "00000000")


if __name__ == "__main__":
    unittest.main(verbosity=2)
