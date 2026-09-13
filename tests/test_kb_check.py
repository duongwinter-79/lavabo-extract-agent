"""What the validator has to catch before a price reaches a customer.

Every case here is a way a real spreadsheet goes wrong, not a way a parser goes wrong:
the shop writes 2tr850 because that is how the whole trade writes it, leaves a promo
without an end date because the promo is "obviously" over, and re-saves a file without
re-checking a single row. None of those look like mistakes inside Excel. All of them
become a wrong number quoted by the Page.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook, load_workbook                      # noqa: E402

from lavabo.kb import spec                                        # noqa: E402
from lavabo.kb.check import check_intake                          # noqa: E402
from lavabo.kb.templates import write_intake                      # noqa: E402

TODAY = date.today()

GOOD = {
    "ma_sp": "BC52-80-TRANG",
    "ten_sp": "Tủ lavabo BC52 80cm",
    "loai": "tủ lavabo",
    "kich_thuoc": "80cm",
    "chat_lieu": "nhựa PVC",
    "mau": "trắng",
    "don_vi": "bộ",
    "gia_niem_yet": 2850000,
    "gia_km": None,
    "km_den_ngay": None,
    "tinh_trang": "đặt trước",
    "thoi_gian_giao": "3-5 ngày",
    "bao_hanh": "12 tháng",
    "gia_gom": "chưa gồm ship",
    "anh": "BC52-80-TRANG__front.jpg",
    "ghi_chu_tu_van": "Đã kèm chậu và vòi.",
    "cap_nhat_ngay": TODAY,
}


def write_catalog(directory: Path, rows: list[dict], *, columns: list[str] | None = None) -> None:
    columns = columns or spec.CATALOG.names
    wb = Workbook()
    ws = wb.active
    ws.title = "Dữ liệu"
    ws.append(columns)
    for row in rows:
        ws.append([row.get(c) for c in columns])
    wb.save(directory / "catalog.xlsx")


class CatalogCase(unittest.TestCase):
    def check(self, rows: list[dict], *, columns: list[str] | None = None):
        self.dir = Path(tempfile.mkdtemp())
        write_catalog(self.dir, rows, columns=columns)
        problems = check_intake(self.dir)
        return [p for p in problems if p.file == "catalog.xlsx"]

    def fatal(self, rows, **kw) -> list[str]:
        return [str(p) for p in self.check(rows, **kw) if p.fatal]

    def warnings(self, rows, **kw) -> list[str]:
        return [str(p) for p in self.check(rows, **kw) if not p.fatal]


class AGoodCatalogPasses(CatalogCase):
    def test_no_fatal_problems(self):
        self.assertEqual(self.fatal([GOOD]), [])

    def test_a_blank_optional_column_is_only_a_warning(self):
        row = {**GOOD, "ghi_chu_tu_van": None, "anh": None}
        self.assertEqual(self.fatal([row]), [])
        self.assertEqual(len(self.warnings([row])), 2)


class PricesThatAreNotNumbers(CatalogCase):
    def test_vietnamese_shorthand_is_rejected_but_understood(self):
        found = self.fatal([{**GOOD, "gia_niem_yet": "2tr850"}])
        self.assertEqual(len(found), 1)
        # Rejecting is not enough: the shop has to be told what to type instead, and
        # money.parse_vnd already knows. Silently accepting it is the actual danger --
        # a column that is sometimes 2850000 and sometimes "2tr850" breaks the feed.
        self.assertIn("2850000", found[0])

    def test_a_range_is_not_a_price(self):
        for written in ("2-3tr", "2 - 3 triệu", "từ 2tr850"):
            with self.subTest(written):
                found = self.fatal([{**GOOD, "gia_niem_yet": written}])
                self.assertEqual(len(found), 1)
                self.assertIn("khoảng giá", found[0])

    def test_a_formatted_number_typed_as_text_says_what_to_type(self):
        found = self.fatal([{**GOOD, "gia_niem_yet": "2.850.000"}])
        self.assertEqual(len(found), 1)
        self.assertIn("2850000", found[0])

    def test_a_real_integer_passes(self):
        self.assertEqual(self.fatal([{**GOOD, "gia_niem_yet": 2850000}]), [])


class Promotions(CatalogCase):
    def test_a_promo_price_without_an_end_date_is_fatal(self):
        found = self.fatal([{**GOOD, "gia_km": 2550000, "km_den_ngay": None}])
        self.assertEqual(len(found), 1)
        self.assertIn("km_den_ngay", found[0])

    def test_an_expired_promo_still_in_the_file_is_fatal(self):
        row = {**GOOD, "gia_km": 2550000, "km_den_ngay": TODAY - timedelta(days=3)}
        found = self.fatal([row])
        self.assertTrue(any("hết hạn" in f for f in found), found)

    def test_a_promo_price_above_the_list_price_is_fatal(self):
        row = {**GOOD, "gia_km": 3000000, "km_den_ngay": TODAY + timedelta(days=10)}
        self.assertTrue(any("không nhỏ hơn" in f for f in self.fatal([row])))

    def test_a_live_promo_passes(self):
        row = {**GOOD, "gia_km": 2550000, "km_den_ngay": TODAY + timedelta(days=30)}
        self.assertEqual(self.fatal([row]), [])


class Identity(CatalogCase):
    def test_a_repeated_ma_sp_is_fatal(self):
        found = self.fatal([GOOD, {**GOOD, "ten_sp": "Tủ khác"}])
        self.assertTrue(any("đã có ở dòng 2" in f for f in found), found)

    def test_a_repeated_name_is_only_a_warning(self):
        found = self.warnings([GOOD, {**GOOD, "ma_sp": "BC52-80-XAM"}])
        self.assertTrue(any("trùng" in f for f in found), found)

    def test_the_example_row_left_in_place_is_flagged(self):
        row = {**GOOD, "ma_sp": f"{spec.EXAMPLE_MARKER}-BC52"}
        self.assertTrue(any("ví dụ" in w for w in self.warnings([row])))


class DeliveryPromises(CatalogCase):
    def test_a_named_day_is_flagged(self):
        for written in ("thứ 5", "ngày 20", "20/9"):
            with self.subTest(written):
                found = self.warnings([{**GOOD, "thoi_gian_giao": written}])
                self.assertTrue(any("ngày cụ thể" in w for w in found), found)

    def test_a_range_is_fine(self):
        found = self.warnings([{**GOOD, "thoi_gian_giao": "3-5 ngày"}])
        self.assertFalse(any("ngày cụ thể" in w for w in found), found)


class Freshness(CatalogCase):
    def test_a_row_nobody_has_checked_this_quarter_is_fatal(self):
        row = {**GOOD, "cap_nhat_ngay": TODAY - timedelta(days=spec.STALE_FAIL_DAYS + 1)}
        self.assertTrue(any("chưa được kiểm tra" in f for f in self.fatal([row])))

    def test_a_month_old_row_is_only_a_warning(self):
        row = {**GOOD, "cap_nhat_ngay": TODAY - timedelta(days=spec.STALE_WARN_DAYS + 5)}
        self.assertEqual(self.fatal([row]), [])
        self.assertTrue(any("chưa kiểm tra lại" in w for w in self.warnings([row])))


class Shape(CatalogCase):
    def test_a_missing_required_column_stops_the_file(self):
        columns = [c for c in spec.CATALOG.names if c != "gia_niem_yet"]
        found = self.fatal([GOOD], columns=columns)
        self.assertEqual(len(found), 1)
        self.assertIn("gia_niem_yet", found[0])

    def test_column_order_does_not_matter(self):
        self.assertEqual(self.fatal([GOOD], columns=list(reversed(spec.CATALOG.names))), [])

    def test_an_unknown_tinh_trang_is_fatal(self):
        found = self.fatal([{**GOOD, "tinh_trang": "sắp về"}])
        self.assertEqual(len(found), 1)
        self.assertIn("đặt trước", found[0])

    def test_a_blank_required_cell_is_fatal(self):
        self.assertTrue(any("đang trống" in f
                            for f in self.fatal([{**GOOD, "don_vi": None}])))


class TheWholePack(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp()) / "intake"

    def test_a_freshly_generated_pack_has_no_spreadsheet_errors(self):
        """init -> check is the loop the shop runs. If the blank workbooks cannot pass
        the validator, the spec and the form have already drifted apart."""
        write_intake(self.dir)
        problems = check_intake(self.dir)
        sheets = [str(p) for p in problems if p.fatal and p.file.endswith(".xlsx")]
        self.assertEqual(sheets, [])
        self.assertTrue(any("ví dụ" in str(p) for p in problems))

    def test_a_freshly_generated_pack_does_not_pass_overall(self):
        """An empty pack is not a usable pack. Every markdown form comes back saying
        which answers are still missing, which is the whole point of the marker."""
        write_intake(self.dir)
        fatal = [p for p in check_intake(self.dir) if p.fatal]
        self.assertEqual({p.file for p in fatal},
                         {d.filename for d in spec.DOC_SPECS})

    def test_init_never_overwrites_a_filled_in_file(self):
        write_intake(self.dir)
        catalog = self.dir / "catalog.xlsx"
        write_catalog(self.dir, [{**GOOD, "ten_sp": "Cái shop đã điền"}])
        written, skipped = write_intake(self.dir)
        self.assertEqual(written, [])
        self.assertIn(catalog, skipped)
        ws = load_workbook(catalog)["Dữ liệu"]
        self.assertEqual(ws.cell(row=2, column=2).value, "Cái shop đã điền")

    def test_force_overwrites(self):
        write_intake(self.dir)
        write_catalog(self.dir, [{**GOOD, "ten_sp": "Cái shop đã điền"}])
        written, skipped = write_intake(self.dir, force=True)
        self.assertEqual(skipped, [])
        self.assertEqual(len(written),
                         len(spec.SHEETS) + len(spec.DOC_SPECS) + len(spec.STATIC_DOCS))

    def test_a_missing_required_file_is_reported(self):
        write_intake(self.dir)
        (self.dir / "shipping.xlsx").unlink()
        problems = check_intake(self.dir)
        self.assertTrue(any(p.file == "shipping.xlsx" and p.fatal for p in problems))

    def test_an_optional_file_may_be_absent(self):
        write_intake(self.dir)
        (self.dir / "synonyms.xlsx").unlink()
        problems = check_intake(self.dir)
        self.assertFalse(any(p.file == "synonyms.xlsx" for p in problems))

    def test_a_price_in_the_faq_is_a_warning(self):
        """docs/13 §4: the catalogue is the only place a number lives. Two copies means
        one stale copy, and the model reads whichever it likes."""
        write_intake(self.dir)
        wb = Workbook()
        ws = wb.active
        ws.title = "Dữ liệu"
        ws.append(spec.FAQ.names)
        ws.append(["Tủ 80 bao nhiêu ạ?", "Dạ bên em 2.850.000đ ạ", "sản phẩm"])
        wb.save(self.dir / "faq.xlsx")
        problems = [p for p in check_intake(self.dir) if p.file == "faq.xlsx"]
        self.assertEqual(len(problems), 1)
        self.assertFalse(problems[0].fatal)
        self.assertIn("catalog.xlsx", str(problems[0]))

    def test_an_image_mapped_to_an_unknown_product_does_not_block_the_upload(self):
        """Found by running `kb feed` on a real catalogue: the shop had emptied
        catalog.xlsx of example rows and left images.xlsx untouched, and the pack
        hard-failed over a photo mapping. A mapping that points nowhere costs a missing
        photo, not a wrong price — the line docs/13 draws for what is fatal."""
        write_intake(self.dir)
        write_catalog(self.dir, [GOOD])          # no VI-DU rows left
        problems = [p for p in check_intake(self.dir) if p.file == "images.xlsx"]
        self.assertEqual(len(problems), 1)
        self.assertFalse(problems[0].fatal)
        self.assertIn("ví dụ", str(problems[0]))

    def test_a_typo_in_an_image_mapping_says_what_it_costs(self):
        write_intake(self.dir)
        write_catalog(self.dir, [GOOD])
        wb = Workbook()
        ws = wb.active
        ws.title = "Dữ liệu"
        ws.append(spec.IMAGES.names)
        ws.append(["a.jpg", "BC52-80-TRNAG", "front", ""])
        wb.save(self.dir / "images.xlsx")
        problems = [p for p in check_intake(self.dir) if p.file == "images.xlsx"]
        self.assertEqual(len(problems), 1)
        self.assertFalse(problems[0].fatal)
        self.assertIn("sẽ không bao giờ được gửi", str(problems[0]))

    def test_the_naming_instructions_are_not_mistaken_for_a_photo(self):
        """kb init drops a .txt in images/ explaining the naming rule. Counting it as an
        unattached photo told the shop to fix something we put there ourselves."""
        write_intake(self.dir)
        problems = [str(p) for p in check_intake(self.dir)]
        self.assertFalse(any("không gắn với sản phẩm" in p for p in problems), problems)

    def test_a_real_unattached_photo_is_still_flagged(self):
        """IMG_4821.jpg would not do: the images.xlsx example row already maps it to a
        product, and a mapped photo is an attached photo."""
        write_intake(self.dir)
        (self.dir / "images" / "IMG_9999.jpg").write_bytes(b"jpeg-ish")
        problems = [str(p) for p in check_intake(self.dir)]
        self.assertTrue(any("không gắn với sản phẩm" in p for p in problems), problems)

    def test_an_image_named_in_the_catalog_but_missing_from_disk_is_flagged(self):
        write_intake(self.dir)
        (self.dir / "images" / "BC52-80-TRANG__front.jpg").write_bytes(b"not really a jpeg")
        write_catalog(self.dir, [{**GOOD, "anh": "BC52-80-TRANG__front.jpg; khong-co.jpg"}])
        problems = [str(p) for p in check_intake(self.dir) if not p.fatal]
        self.assertTrue(any("khong-co.jpg" in p for p in problems), problems)
        self.assertFalse(any("__front.jpg" in p for p in problems), problems)


if __name__ == "__main__":
    unittest.main()


class TheMarkdownForms(unittest.TestCase):
    """The blanks nobody filled in.

    A file returned untouched and a file somebody deliberately left empty look identical
    without a marker, so the shop's "we sent you everything" and our "half of it is
    blank" were previously both true and unarguable.
    """

    def setUp(self):
        self.dir = Path(tempfile.mkdtemp()) / "intake"
        write_intake(self.dir)

    def problems(self, filename="store.md"):
        return [p for p in check_intake(self.dir) if p.file == filename]

    def fill(self, filename, answer="xong"):
        path = self.dir / filename
        path.write_text(path.read_text(encoding="utf-8").replace(spec.PLACEHOLDER, answer),
                        encoding="utf-8")

    def test_an_untouched_form_names_every_missing_answer(self):
        fatal = [p for p in self.problems() if p.fatal]
        self.assertEqual(len(fatal), 1)
        self.assertIn("Tên shop", str(fatal[0]))
        self.assertIn("Hotline", str(fatal[0]))

    def test_optional_blanks_do_not_block(self):
        warned = [p for p in self.problems() if not p.fatal]
        self.assertEqual(len(warned), 1)
        self.assertIn("Website", str(warned[0]))
        self.assertNotIn("Hotline", str(warned[0]))

    def test_a_filled_form_is_clean(self):
        self.fill("store.md")
        self.assertEqual(self.problems(), [])

    def test_the_instructions_are_not_mistaken_for_a_blank(self):
        """The header explains what the marker means, so it contains one. A shop fills in
        the fields and leaves that paragraph alone -- as they should -- and the form must
        still come back clean instead of reporting one missing answer forever."""
        path = self.dir / "store.md"
        filled = [ln.replace(spec.PLACEHOLDER, "xong") if ln.startswith("-") else ln
                  for ln in path.read_text(encoding="utf-8").splitlines()]
        path.write_text("\n".join(filled), encoding="utf-8")

        self.assertIn(spec.PLACEHOLDER, path.read_text(encoding="utf-8"))
        self.assertEqual(self.problems(), [])

    def test_reordering_and_deleting_sections_is_allowed(self):
        """They will rearrange it. Matching by label rather than position means an edit
        that loses nothing does not read as a missing answer."""
        path = self.dir / "store.md"
        lines = path.read_text(encoding="utf-8").splitlines()
        kept = [ln for ln in lines if "Địa chỉ 2" not in ln and "Ngày nghỉ" not in ln]
        path.write_text("\n".join(reversed(kept)), encoding="utf-8")
        fatal = [p for p in self.problems() if p.fatal]
        self.assertEqual(len(fatal), 1)
        self.assertIn("Tên shop", str(fatal[0]))

    def test_a_sample_conversation_is_tracked_by_its_heading(self):
        problems = [str(p) for p in self.problems("voice.md") if p.fatal]
        self.assertTrue(any("Hội thoại mẫu 1" in p for p in problems), problems)

    def test_writing_the_samples_clears_them(self):
        self.fill("voice.md", "Khách: tủ 80 bao nhiêu ạ?\nShop: dạ bên em...")
        self.assertEqual(self.problems("voice.md"), [])

    def test_a_deleted_form_is_reported(self):
        (self.dir / "policies.md").unlink()
        problems = self.problems("policies.md")
        self.assertEqual(len(problems), 1)
        self.assertIn("thiếu file", str(problems[0]))
