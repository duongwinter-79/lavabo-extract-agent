"""Naming photos nobody named.

Zalo delivers `received_8837221.jpg` for a product the sender never mentioned. Asking the
shop about each photo is forty messages nobody answers; asking once about a numbered sheet
of their own pictures is one message and one reply. These tests are mostly about the sheet
being unambiguous, because its single job is "which number is this photo".
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook                                # noqa: E402
from PIL import Image                                             # noqa: E402

from lavabo.kb.contact import MAP_FILENAME, build, read_mapping   # noqa: E402
from lavabo.kb.media import organise                              # noqa: E402


def photo(path: Path, size=(1600, 1200)) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (180, 190, 200)).save(path, "JPEG")


class TheSheet(unittest.TestCase):
    def setUp(self):
        self.source = Path(tempfile.mkdtemp()) / "zalo"
        self.out = Path(tempfile.mkdtemp()) / "cs"
        for name in ("received_1.jpg", "received_2.jpg", "z5123991.jpg"):
            photo(self.source / name)

    def test_it_numbers_every_photo_and_writes_the_mapping(self):
        sheet = build(self.source, self.out)
        self.assertEqual(sheet.count, 3)
        self.assertEqual(len(sheet.pages), 1)
        self.assertTrue(sheet.pages[0].exists())
        self.assertEqual(sheet.mapping.name, MAP_FILENAME)

    def test_the_mapping_lists_filenames_in_the_numbered_order(self):
        sheet = build(self.source, self.out)
        rows = list(load_workbook(sheet.mapping)["Dữ liệu"].iter_rows(min_row=2,
                                                                     values_only=True))
        self.assertEqual([r[0] for r in rows], [1, 2, 3])
        self.assertEqual([r[1] for r in rows],
                         ["received_1.jpg", "received_2.jpg", "z5123991.jpg"])
        self.assertTrue(all(r[2] is None or r[2] == "" for r in rows))

    def test_many_photos_spill_onto_more_pages(self):
        for n in range(4, 20):
            photo(self.source / f"p{n}.jpg", size=(800, 600))
        sheet = build(self.source, self.out)
        self.assertEqual(sheet.count, 19)
        self.assertGreater(len(sheet.pages), 1)

    def test_a_corrupt_photo_does_not_take_the_sheet_down(self):
        (self.source / "broken.jpg").write_bytes(b"nope")
        sheet = build(self.source, self.out)
        self.assertEqual(sheet.count, 4)
        self.assertTrue(sheet.pages[0].exists())

    def test_an_empty_folder_says_so(self):
        with self.assertRaises(ValueError):
            build(Path(tempfile.mkdtemp()), self.out)


class ApplyingTheAnswer(unittest.TestCase):
    def setUp(self):
        self.source = Path(tempfile.mkdtemp()) / "zalo"
        self.out = Path(tempfile.mkdtemp()) / "cs"
        self.images = Path(tempfile.mkdtemp()) / "images"
        for name in ("a.jpg", "b.jpg", "c.jpg"):
            photo(self.source / name)
        self.sheet = build(self.source, self.out)

    def answer(self, rows: dict[int, tuple[str, str]]):
        wb = load_workbook(self.sheet.mapping)
        ws = wb["Dữ liệu"]
        for row in ws.iter_rows(min_row=2):
            if got := rows.get(row[0].value):
                row[2].value, row[3].value = got
        wb.save(self.sheet.mapping)

    def test_the_shops_reply_becomes_product_filenames(self):
        self.answer({1: ("BC52-80", "front"), 2: ("BC52-80", "lapdat"),
                     3: ("GUONG-60", "front")})
        report = organise(self.source, self.images,
                          mapping=read_mapping(self.sheet.mapping))
        self.assertEqual(report.photo_count, 3)
        self.assertEqual({p.name for p in self.images.iterdir()},
                         {"BC52-80__front.jpg", "BC52-80__lapdat.jpg",
                          "GUONG-60__front.jpg"})

    def test_a_row_left_blank_is_reported_never_guessed(self):
        """The premise of this whole route is that nobody could tell which product it
        was. Filing it under a plausible neighbour would be the one unforgivable move."""
        self.answer({1: ("BC52-80", "front")})
        report = organise(self.source, self.images,
                          mapping=read_mapping(self.sheet.mapping))
        self.assertEqual(report.photo_count, 1)
        self.assertEqual(sorted(report.unmapped), ["b.jpg", "c.jpg"])

    def test_two_photos_for_one_product_get_different_names(self):
        self.answer({1: ("BC52-80", ""), 2: ("BC52-80", ""), 3: ("BC52-80", "")})
        report = organise(self.source, self.images,
                          mapping=read_mapping(self.sheet.mapping))
        self.assertEqual(len(set(report.products["BC52-80"])), 3)

    def test_a_mapping_without_the_needed_columns_is_rejected(self):
        from openpyxl import Workbook

        broken = self.out / "broken.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Dữ liệu"
        ws.append(["so_thu_tu", "ghi_chu"])
        wb.save(broken)
        with self.assertRaises(ValueError):
            read_mapping(broken)


if __name__ == "__main__":
    unittest.main()
