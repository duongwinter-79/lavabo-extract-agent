"""Phone media into the pack.

The shop's photos and demo videos live on a phone, where renaming files is not a thing
anyone does. So the contract is: a folder per mã SP, any filenames inside, and this
produces the names the pack expects. The cases below are the ways a real phone dump is
messy -- HEIC from an iPhone, a loose photo nobody filed, two shots that both look like
the front, a video where only some frames are in focus.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from PIL import Image                                             # noqa: E402

from lavabo.kb.media import (MAX_EDGE, organise, organise_workbook,  # noqa: E402
                             write_mapping)
from lavabo.video import ffmpeg_path                              # noqa: E402


def photo(path: Path, size=(3024, 4032)) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (180, 190, 200)).save(path, "JPEG", quality=90)


class Sorting(unittest.TestCase):
    def setUp(self):
        self.source = Path(tempfile.mkdtemp()) / "phone"
        self.images = Path(tempfile.mkdtemp()) / "images"

    def run_it(self, **kw):
        return organise(self.source, self.images, **kw)

    def names(self) -> set[str]:
        return {p.name for p in self.images.iterdir()}

    def test_a_folder_per_product_becomes_named_photos(self):
        photo(self.source / "BC52-80" / "IMG_4821.jpg")
        photo(self.source / "GUONG-60" / "IMG_5001.jpg")
        report = self.run_it()
        self.assertEqual(report.photo_count, 2)
        self.assertEqual(self.names(), {"BC52-80__front.jpg", "GUONG-60__front.jpg"})

    def test_a_filename_hinting_a_role_is_believed(self):
        photo(self.source / "BC52-80" / "anh-lapdat.jpg")
        self.run_it()
        self.assertIn("BC52-80__lapdat.jpg", self.names())

    def test_two_photos_hinting_the_same_role_do_not_overwrite_each_other(self):
        """The bug this exists for: both matched "front", one silently replaced the
        other on disk, and the report still claimed two photos."""
        photo(self.source / "BC52-80" / "front-1.jpg")
        photo(self.source / "BC52-80" / "anh-chinh.jpg")
        report = self.run_it()
        self.assertEqual(report.photo_count, 2)
        self.assertEqual(len(self.names()), 2)

    def test_photos_are_downscaled(self):
        photo(self.source / "BC52-80" / "IMG_4821.jpg")
        self.run_it()
        with Image.open(self.images / "BC52-80__front.jpg") as img:
            self.assertLessEqual(max(img.size), MAX_EDGE)

    def test_a_sideways_photo_is_not_left_sideways(self):
        path = self.source / "BC52-80" / "IMG_4821.jpg"
        photo(path, size=(1200, 800))
        with Image.open(path) as img:
            img.save(path, "JPEG", exif=Image.Exif())
        self.run_it()
        self.assertIn("BC52-80__front.jpg", self.names())

    def test_heic_is_reported_rather_than_silently_dropped(self):
        (self.source / "BC52-80").mkdir(parents=True)
        (self.source / "BC52-80" / "IMG_9000.HEIC").write_bytes(b"not really heic")
        report = self.run_it()
        self.assertEqual(report.heic, ["BC52-80/IMG_9000.HEIC"])
        self.assertEqual(report.photo_count, 0)

    def test_a_loose_photo_is_reported_never_guessed_at(self):
        """An unattached photo costs a missing picture. A misattached one puts the wrong
        product in front of a customer."""
        photo(self.source / "khong-biet.jpg")
        report = self.run_it()
        self.assertEqual(report.unmapped, ["khong-biet.jpg"])
        self.assertEqual(report.photo_count, 0)

    def test_a_loose_photo_named_after_its_product_is_accepted(self):
        photo(self.source / "BC52-80__front.jpg")
        report = self.run_it()
        self.assertEqual(report.photo_count, 1)
        self.assertIn("BC52-80__front.jpg", self.names())

    def test_an_empty_product_folder_is_visible_in_the_report(self):
        (self.source / "SEN-CAY").mkdir(parents=True)
        report = self.run_it()
        self.assertEqual(report.products["SEN-CAY"], [])

    def test_a_code_not_in_the_catalogue_is_flagged(self):
        photo(self.source / "GO-NHAM" / "a.jpg")
        report = self.run_it(known_skus={"bc52-80"})
        self.assertEqual(report.unknown_sku, ["GO-NHAM"])

    def test_a_corrupt_file_is_skipped_not_fatal(self):
        (self.source / "BC52-80").mkdir(parents=True)
        (self.source / "BC52-80" / "broken.jpg").write_bytes(b"nope")
        report = self.run_it()
        self.assertEqual(report.photo_count, 0)
        self.assertEqual(report.skipped, ["broken.jpg"])


class DemoVideos(unittest.TestCase):
    def setUp(self):
        self.source = Path(tempfile.mkdtemp()) / "phone"
        self.images = Path(tempfile.mkdtemp()) / "images"
        (self.source / "SEN-CAY").mkdir(parents=True)
        subprocess.run(
            [ffmpeg_path(), "-loglevel", "error", "-f", "lavfi",
             # Six seconds: frames are sampled at 1fps, so a three-second clip cannot
             # answer a request for four stills and the test would be measuring ffmpeg.
             "-i", "testsrc=size=320x240:rate=10", "-t", "6", "-pix_fmt", "yuv420p",
             str(self.source / "SEN-CAY" / "demo.mp4")], check=True)

    def test_stills_are_pulled_out_of_the_video(self):
        report = organise(self.source, self.images, frames=3)
        self.assertEqual(report.videos, ["SEN-CAY/demo.mp4"])
        self.assertEqual(len(report.products["SEN-CAY"]), 3)
        self.assertEqual(len({p.name for p in self.images.iterdir()}), 3)

    def test_every_requested_frame_survives(self):
        """Frame names carry no role information, so hinting off them made two frames
        collide on "front" and one was lost."""
        report = organise(self.source, self.images, frames=4)
        self.assertEqual(len(set(report.products["SEN-CAY"])), 4)


class TheMapping(unittest.TestCase):
    def test_it_records_which_product_each_renamed_photo_belongs_to(self):
        from openpyxl import load_workbook

        source = Path(tempfile.mkdtemp()) / "phone"
        images = Path(tempfile.mkdtemp()) / "images"
        photo(source / "BC52-80" / "IMG_1.jpg")
        photo(source / "BC52-80" / "chi-tiet.jpg")
        report = organise(source, images)

        out = images.parent / "images.xlsx"
        write_mapping(report, out)
        rows = list(load_workbook(out)["Dữ liệu"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r[1] == "BC52-80" for r in rows), rows)
        self.assertIn("detail", [r[2] for r in rows])


if __name__ == "__main__":
    unittest.main()


class PhotosPastedIntoTheWorkbook(unittest.TestCase):
    """The shop will paste pictures next to the products, because that is the obvious
    thing to do. openpyxl keeps the anchor, so the row identifies the product and the
    intuitive act stops being a dead end."""

    def setUp(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage

        self.dir = Path(tempfile.mkdtemp())
        self.images = self.dir / "images"
        for n in range(1, 4):
            photo(self.dir / f"p{n}.jpg", size=(900, 700))

        wb = Workbook()
        ws = wb.active
        ws.title = "Dữ liệu"
        ws.append(["ma_sp", "ten_sp"])
        for code in ("BC52-80", "GUONG-60", "SEN-CAY"):
            ws.append([code, f"Sản phẩm {code}"])
        for n, row in enumerate([2, 3, 4], 1):
            image = XLImage(str(self.dir / f"p{n}.jpg"))
            image.width, image.height = 80, 60
            ws.add_image(image, f"C{row}")
        self.book = self.dir / "pasted.xlsx"
        wb.save(self.book)

    def test_each_image_lands_under_the_product_on_its_row(self):
        report = organise_workbook(self.book, self.images)
        self.assertEqual(report.photo_count, 3)
        self.assertEqual({p.name for p in self.images.iterdir()},
                         {"BC52-80__front.jpg", "GUONG-60__front.jpg", "SEN-CAY__front.jpg"})

    def test_extracted_images_are_downscaled_like_any_other(self):
        organise_workbook(self.book, self.images)
        with Image.open(self.images / "BC52-80__front.jpg") as img:
            self.assertLessEqual(max(img.size), MAX_EDGE)

    def test_two_photos_on_one_row_both_survive(self):
        from openpyxl import load_workbook
        from openpyxl.drawing.image import Image as XLImage

        wb = load_workbook(self.book)
        ws = wb["Dữ liệu"]
        extra = XLImage(str(self.dir / "p1.jpg"))
        extra.width, extra.height = 80, 60
        ws.add_image(extra, "D2")
        wb.save(self.book)

        report = organise_workbook(self.book, self.images)
        self.assertEqual(len(report.products["BC52-80"]), 2)

    def test_an_image_floating_off_the_products_is_reported_not_guessed(self):
        from openpyxl import load_workbook
        from openpyxl.drawing.image import Image as XLImage

        wb = load_workbook(self.book)
        ws = wb["Dữ liệu"]
        stray = XLImage(str(self.dir / "p1.jpg"))
        stray.width, stray.height = 80, 60
        ws.add_image(stray, "C40")
        wb.save(self.book)

        report = organise_workbook(self.book, self.images)
        self.assertEqual(len(report.unmapped), 1)
        self.assertEqual(report.photo_count, 3)

    def test_a_code_not_in_the_catalogue_is_still_flagged(self):
        report = organise_workbook(self.book, self.images, known_skus={"bc52-80"})
        self.assertEqual(report.unknown_sku, ["GUONG-60", "SEN-CAY"])
