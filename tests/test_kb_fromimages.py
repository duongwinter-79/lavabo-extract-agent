"""Reading a catalogue off the pictures, without inventing one.

The shop's marketing images carry kích thước, màu and often a price. That is a catalogue
nobody typed up. It is also a trap: a number printed on a promotional image is what
somebody charged on some day, possibly a different shop, and the moment it looks like a
price list somebody will treat it as one. These tests are mostly about that distinction
surviving.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook                                # noqa: E402
from PIL import Image                                             # noqa: E402

from lavabo.kb.fromimages import (DRAFT_COLUMNS, SCHEMA, read_folder,  # noqa: E402
                                  write_draft)


class Completion:
    def __init__(self, values, input_tokens=10, output_tokens=5):
        self.values = values
        self.input_tokens = input_tokens
        self.output_tokens = output_tokens


class StubExtractor:
    """Stands in for a real provider, so the pipeline is testable without an API key."""

    def __init__(self, answers=None, fail_on=()):
        self.answers = answers or {}
        self.fail_on = set(fail_on)
        self.calls: list[int] = []

    def complete_json_images(self, system, user, images, schema, *, mime_type="image/png",
                             max_tokens=0):
        self.calls.append(len(images))
        blob = images[0]
        if blob in self.fail_on:
            raise RuntimeError("model exploded")
        return Completion(self.answers.get(blob, _blank()))


def _blank():
    return {key: None for key in SCHEMA["properties"]} | {"la_san_pham": False}


def photo(path: Path, colour=(200, 200, 200)) -> bytes:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (600, 400), colour).save(path, "JPEG")
    return path.read_bytes()


class ReadingTheFolder(unittest.TestCase):
    def setUp(self):
        self.source = Path(tempfile.mkdtemp()) / "anh"
        self.a = photo(self.source / "a.jpg", (200, 60, 60))
        self.b = photo(self.source / "b.jpg", (60, 200, 60))

    def test_one_call_per_image(self):
        """A bad image cannot poison a batch, and a re-run costs only what failed."""
        stub = StubExtractor()
        read_folder(self.source, stub)
        self.assertEqual(stub.calls, [1, 1])

    def test_it_counts_products_and_prices_separately(self):
        stub = StubExtractor({
            self.a: _blank() | {"la_san_pham": True, "ten_sp": "Tủ 80", "gia": 2850000},
            self.b: _blank() | {"la_san_pham": True, "ten_sp": "Gương"},
        })
        report = read_folder(self.source, stub)
        self.assertEqual(len(report.products), 2)
        self.assertEqual(len(report.with_price), 1)

    def test_a_non_product_image_is_not_counted_as_one(self):
        stub = StubExtractor({self.a: _blank() | {"la_san_pham": True},
                              self.b: _blank()})
        self.assertEqual(len(read_folder(self.source, stub).products), 1)

    def test_one_failure_does_not_stop_the_run(self):
        stub = StubExtractor(fail_on=[self.a])
        report = read_folder(self.source, stub)
        self.assertEqual(len(report.readings), 2)
        self.assertEqual(len(report.failed), 1)
        self.assertEqual(report.failed[0].path.name, "a.jpg")

    def test_limit_stops_early(self):
        stub = StubExtractor()
        read_folder(self.source, stub, limit=1)
        self.assertEqual(stub.calls, [1])

    def test_tokens_are_totalled(self):
        report = read_folder(self.source, StubExtractor())
        self.assertEqual(report.input_tokens, 20)
        self.assertEqual(report.output_tokens, 10)


class TheDraft(unittest.TestCase):
    def setUp(self):
        self.source = Path(tempfile.mkdtemp()) / "anh"
        self.a = photo(self.source / "a.jpg")
        self.target = Path(tempfile.mkdtemp()) / "draft.xlsx"

    def draft(self, answers=None, fail_on=()):
        report = read_folder(self.source, StubExtractor(answers, fail_on))
        write_draft(report, self.target)
        return list(load_workbook(self.target)["Dữ liệu"].iter_rows(values_only=True))

    def test_the_first_row_says_it_is_not_a_price_list(self):
        """Anything shaped like a spreadsheet of prices gets treated as one."""
        rows = self.draft()
        self.assertIn("BẢN NHÁP", str(rows[0][0]))
        self.assertIn("có thể cũ", str(rows[0][0]))

    def test_the_price_column_is_named_for_what_it_is(self):
        self.assertIn("gia_doc_duoc", DRAFT_COLUMNS)
        self.assertNotIn("gia_niem_yet", DRAFT_COLUMNS)

    def test_it_records_where_the_number_was_seen(self):
        """Confirming a price should mean looking at one picture, not trusting a cell."""
        rows = self.draft({self.a: _blank() | {"la_san_pham": True, "gia": 2850000,
                                               "gia_nhin_thay_o": "Giá: 2.850.000đ"}})
        self.assertIn("Giá: 2.850.000đ", rows[2])
        self.assertIn("a.jpg", rows[2])

    def test_it_leaves_the_code_and_the_confirmation_blank(self):
        rows = self.draft({self.a: _blank() | {"la_san_pham": True, "ten_sp": "Tủ 80"}})
        header = list(rows[1])
        row = list(rows[2])
        # openpyxl reads a blank cell back as None, whatever was written.
        self.assertFalse(row[header.index("ma_sp")])
        self.assertFalse(row[header.index("da_kiem_tra")])

    def test_a_failed_image_is_visible_in_the_draft(self):
        rows = self.draft(fail_on=[self.a])
        self.assertEqual(rows[2][1], "LỖI")

    def test_the_schema_never_asks_for_a_product_code(self):
        """A photograph does not carry one, and a model asked for it will invent it."""
        self.assertNotIn("ma_sp", SCHEMA["properties"])


if __name__ == "__main__":
    unittest.main()
