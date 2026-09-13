"""The committed pack must be what `kb init` produces today.

`templates/intake/` exists so whoever forwards the pack to the shop does not need a
Python environment to get it. That convenience is also a liability: it is generated
output living in git, and generated output in git rots quietly. Nobody re-reads a binary
workbook in a diff, so a spec change would leave the committed copy a month out of date
and the shop would fill in the wrong form.

This compares CONTENT, not bytes. openpyxl stamps a creation time inside every workbook,
so the files differ run to run while saying exactly the same thing.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook                                # noqa: E402

from lavabo.kb.templates import write_intake                      # noqa: E402

COMMITTED = ROOT / "templates" / "intake"


def fingerprint(directory: Path) -> dict[str, object]:
    """Everything about the pack that the shop can see."""
    out: dict[str, object] = {}
    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue
        name = str(path.relative_to(directory))
        if path.suffix == ".xlsx":
            book = load_workbook(path)
            out[name] = {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
                         for sheet in book.worksheets}
        else:
            out[name] = path.read_text(encoding="utf-8")
    return out


class TheCommittedPack(unittest.TestCase):
    def setUp(self):
        self.fresh = Path(tempfile.mkdtemp()) / "intake"
        write_intake(self.fresh)

    def test_it_exists(self):
        self.assertTrue(COMMITTED.is_dir(),
                        "templates/intake/ is missing — regenerate with: "
                        "lavabo kb init --dir templates/intake --zip templates/intake-lavabo.zip")

    def test_it_holds_every_file_kb_init_writes(self):
        self.assertEqual(set(fingerprint(COMMITTED)), set(fingerprint(self.fresh)))

    def test_every_file_still_says_the_same_thing(self):
        fresh, committed = fingerprint(self.fresh), fingerprint(COMMITTED)
        for name in sorted(fresh):
            with self.subTest(name):
                self.assertEqual(
                    committed.get(name), fresh[name],
                    f"{name} in templates/intake/ is stale. Regenerate:\n"
                    "  lavabo kb init --dir templates/intake --force "
                    "--zip templates/intake-lavabo.zip")

    def test_the_zip_is_committed_too(self):
        archive = ROOT / "templates" / "intake-lavabo.zip"
        self.assertTrue(archive.is_file())
        self.assertGreater(archive.stat().st_size, 10_000)


if __name__ == "__main__":
    unittest.main()
